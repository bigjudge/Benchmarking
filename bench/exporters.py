"""Export der Benchmarking-Daten als CSV, Excel und Markdown-Report."""

from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import analytics, schema, store

DEFS = schema.all_kpi_defs()

ACCENT = "2A78D6"
LIGHT = "E8F0FB"
GREY = "898781"


def _kpi_keys(include_derived: bool) -> list[str]:
    return schema.REPORTED_KEYS + (schema.DERIVED_KEYS if include_derived else [])


def _fmt(v):
    if v is None:
        return ""
    return round(v, 4) if isinstance(v, float) else v


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------
def to_csv(data: dict, years: list[str], bank_ids: list[str],
           include_derived: bool = True) -> str:
    mat = analytics.matrix(data, years)
    keys = _kpi_keys(include_derived)
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\n")
    writer.writerow(["Bank", "Jahr", "Kennzahl", "Schlüssel", "Kategorie", "Wert",
                     "Einheit", "Herkunft", "Konfidenz", "Quelle"])
    for bank in data.get("banken", []):
        if bank["id"] not in bank_ids:
            continue
        for year in years:
            cells = (mat.get(bank["id"]) or {}).get(year, {})
            for key in keys:
                cell = cells.get(key) or {}
                if cell.get("value") is None:
                    continue
                d = DEFS[key]
                writer.writerow([
                    bank["name"], year, d["label"], key,
                    schema.KATEGORIEN.get(d["kategorie"], d["kategorie"]),
                    _fmt(cell.get("value")), cell.get("unit", ""),
                    cell.get("herkunft", ""), cell.get("confidence", ""),
                    (cell.get("source") or "").replace("\n", " "),
                ])
    return buf.getvalue()


def to_wide_csv(data: dict, year: str, bank_ids: list[str],
                include_derived: bool = True) -> str:
    mat = analytics.matrix(data, [year])
    keys = _kpi_keys(include_derived)
    banks = [b for b in data.get("banken", []) if b["id"] in bank_ids]
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\n")
    writer.writerow(["Kennzahl", "Einheit", "Kategorie"] + [b["name"] for b in banks])
    for key in keys:
        d = DEFS[key]
        row = [d["label"], d["unit"], schema.KATEGORIEN.get(d["kategorie"], "")]
        for b in banks:
            row.append(_fmt((mat.get(b["id"]) or {}).get(year, {}).get(key, {}).get("value")))
        writer.writerow(row)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def to_xlsx(data: dict, years: list[str], bank_ids: list[str],
            include_derived: bool = True) -> bytes:
    mat = analytics.matrix(data, years)
    keys = _kpi_keys(include_derived)
    banks = [b for b in data.get("banken", []) if b["id"] in bank_ids]
    target = store.target_bank(data)
    target_id = target["id"] if target else None

    wb = Workbook()
    thin = Side(style="thin", color="DDDDDD")
    border = Border(bottom=thin)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=ACCENT)

    # --- Übersicht ---
    ws = wb.active
    ws.title = "Übersicht"
    ws["A1"] = data.get("meta", {}).get("titel", "Bank-Benchmarking")
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = f"Export vom {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws["A3"] = f"Zielbank: {target['name'] if target else '–'}"
    ws["A4"] = f"Banken im Vergleich: {len(banks)} · Jahre: {', '.join(years)}"
    ws["A5"] = ("Abgeleitete Kennzahlen sind berechnet und als solche gekennzeichnet."
                if include_derived else "Nur berichtete Kennzahlen enthalten.")
    for r in range(1, 6):
        ws.cell(row=r, column=1).alignment = Alignment(vertical="center")
    ws.column_dimensions["A"].width = 100

    # --- Matrix je Jahr ---
    for year in years:
        sheet = wb.create_sheet(f"Kennzahlen {year}")
        header = ["Kennzahl", "Einheit", "Kategorie", "Art"] + [b["name"] for b in banks]
        sheet.append(header)
        for col in range(1, len(header) + 1):
            c = sheet.cell(row=1, column=col)
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for key in keys:
            d = DEFS[key]
            row = [d["label"], d["unit"], schema.KATEGORIEN.get(d["kategorie"], ""),
                   "berechnet" if d.get("berechnet") else "berichtet"]
            for b in banks:
                row.append(_fmt((mat.get(b["id"]) or {}).get(year, {})
                                .get(key, {}).get("value")))
            sheet.append(row)
        for idx, b in enumerate(banks):
            if b["id"] == target_id:
                letter = get_column_letter(5 + idx)
                for r in range(1, sheet.max_row + 1):
                    sheet[f"{letter}{r}"].fill = PatternFill("solid", fgColor=LIGHT)
        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 14
        sheet.column_dimensions["C"].width = 26
        sheet.column_dimensions["D"].width = 12
        for idx in range(len(banks)):
            sheet.column_dimensions[get_column_letter(5 + idx)].width = 18
        sheet.freeze_panes = "E2"

    # --- Einzelwerte mit Quelle ---
    detail = wb.create_sheet("Einzelwerte & Quellen")
    detail.append(["Bank", "Jahr", "Kennzahl", "Wert", "Einheit", "Herkunft",
                   "Konfidenz", "Quelle / Formel"])
    for col in range(1, 9):
        c = detail.cell(row=1, column=col)
        c.font = head_font
        c.fill = head_fill
    for b in banks:
        for year in years:
            cells = (mat.get(b["id"]) or {}).get(year, {})
            for key in keys:
                cell = cells.get(key) or {}
                if cell.get("value") is None:
                    continue
                detail.append([b["name"], year, DEFS[key]["label"],
                               _fmt(cell.get("value")), cell.get("unit", ""),
                               cell.get("herkunft", ""), cell.get("confidence"),
                               (cell.get("source") or "")[:300]])
    for width, letter in zip((30, 8, 40, 14, 14, 14, 11, 70), "ABCDEFGH"):
        detail.column_dimensions[letter].width = width
    detail.freeze_panes = "A2"

    # --- Scorecard ---
    if years:
        sc = analytics.scorecard(data, years[0], bank_ids)
        sheet = wb.create_sheet("Scorecard")
        kats = list(schema.KATEGORIEN.keys())
        sheet.append(["Bank", "Gesamt-Perzentil", "Bewertete Kennzahlen"]
                     + [schema.KATEGORIEN[k] for k in kats])
        for col in range(1, 4 + len(kats)):
            c = sheet.cell(row=1, column=col)
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True)
        for row in sc["banken"]:
            sheet.append([row["name"], row["gesamt"], row["abgedeckteKpis"]]
                         + [row["kategorien"].get(k) for k in kats])
        sheet.column_dimensions["A"].width = 32
        for i in range(2, 4 + len(kats)):
            sheet.column_dimensions[get_column_letter(i)].width = 17

    # --- Stärken / Schwächen ---
    if years:
        ins = analytics.insights(data, years[0], bank_ids, include_derived)
        sheet = wb.create_sheet("Stärken & Schwächen")
        sheet.append(["Bewertung", "Kennzahl", "Art", "Zielbank", "Peer-Ø",
                      "Abweichung %", "Perzentil", "Rang"])
        for col in range(1, 9):
            c = sheet.cell(row=1, column=col)
            c.font = head_font
            c.fill = head_fill
        for label, items in (("Stärke", ins["staerken"]), ("Schwäche", ins["schwaechen"])):
            for it in items:
                sheet.append([label, it["label"],
                              "berechnet" if it["berechnet"] else "berichtet",
                              _fmt(it["wert"]), _fmt(it["peerAvg"]),
                              it["abweichungPct"], it["perzentil"], it["rang"]])
        sheet.column_dimensions["A"].width = 12
        sheet.column_dimensions["B"].width = 42
        for letter in "CDEFGH":
            sheet.column_dimensions[letter].width = 15

    # --- Methodik ---
    meth = wb.create_sheet("Methodik")
    meth.append(["Einschränkungen und Hinweise"])
    meth["A1"].font = Font(bold=True, size=12)
    row = 3
    for item in data.get("methodik", {}).get("einschraenkungen", []):
        meth.cell(row=row, column=1, value=f"• {item}").alignment = Alignment(
            wrap_text=True, vertical="top")
        row += 1
    meth.cell(row=row + 1, column=1, value="Verwendete Quelldateien").font = Font(bold=True)
    row += 2
    for f in data.get("methodik", {}).get("verwendeteDateien", []):
        meth.cell(row=row, column=1, value=f"• {f}")
        row += 1
    meth.column_dimensions["A"].width = 140

    # --- Kennzahlendefinitionen ---
    dd = wb.create_sheet("Definitionen")
    dd.append(["Schlüssel", "Kennzahl", "Einheit", "Kategorie", "Art", "Richtung",
               "Beschreibung", "Formel"])
    for col in range(1, 9):
        c = dd.cell(row=1, column=col)
        c.font = head_font
        c.fill = head_fill
    for key in keys:
        d = DEFS[key]
        richtung = {True: "höher = besser", False: "niedriger = besser",
                    None: "neutral"}[d["hoeherIstBesser"]]
        dd.append([key, d["label"], d["unit"],
                   schema.KATEGORIEN.get(d["kategorie"], ""),
                   "berechnet" if d.get("berechnet") else "berichtet", richtung,
                   d.get("beschreibung", ""), d.get("formel", "")])
    for width, letter in zip((26, 40, 14, 28, 12, 20, 80, 45), "ABCDEFGH"):
        dd.column_dimensions[letter].width = width
    dd.freeze_panes = "A2"

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.border = border

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------
# Markdown-Report
# --------------------------------------------------------------------------
def to_markdown(data: dict, year: str, bank_ids: list[str],
                include_derived: bool = True) -> str:
    target = store.target_bank(data)
    ins = analytics.insights(data, year, bank_ids, include_derived)
    sc = analytics.scorecard(data, year, bank_ids)
    mat = analytics.matrix(data, [year])
    banks = [b for b in data.get("banken", []) if b["id"] in bank_ids]

    lines = [
        f"# {data.get('meta', {}).get('titel', 'Bank-Benchmarking')} – Berichtsjahr {year}",
        "",
        f"Zielbank: **{target['name'] if target else '–'}**  ",
        f"Vergleichsgruppe: {len(banks)} Institute  ",
        f"Erstellt: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
        "",
        "## Gesamtpositionierung",
        "",
        "| Bank | Gesamt-Perzentil | Bewertete Kennzahlen |",
        "|---|---:|---:|",
    ]
    for row in sc["banken"]:
        mark = " **(Zielbank)**" if row["istZielbank"] else ""
        lines.append(f"| {row['name']}{mark} | "
                     f"{row['gesamt'] if row['gesamt'] is not None else '–'} | "
                     f"{row['abgedeckteKpis']} |")

    lines += ["", "## Stärken gegenüber dem Peer-Durchschnitt", "",
              "Absolute Grössenkennzahlen wie Bilanzsumme oder verwaltete Kundenvermögen "
              "bleiben hier aussen vor. Sie bilden die Grösse des Instituts ab, nicht "
              "seine Leistung, und würden jedes kleinere Haus pauschal als schwach "
              "erscheinen lassen.", ""]
    if ins["staerken"]:
        lines += ["| Kennzahl | Zielbank | Peer-Ø | Abweichung | Rang |", "|---|---:|---:|---:|---:|"]
        for it in ins["staerken"][:15]:
            art = " *(berechnet)*" if it["berechnet"] else ""
            lines.append(f"| {it['label']}{art} | {_fmt(it['wert'])} {it['unit']} | "
                         f"{_fmt(it['peerAvg'])} | "
                         f"{'+' if (it['abweichungPct'] or 0) >= 0 else ''}"
                         f"{it['abweichungPct']} % | {it['rang'] or '–'} |")
    else:
        lines.append("Keine auswertbaren Kennzahlen über dem Peer-Durchschnitt.")

    lines += ["", "## Schwächen gegenüber dem Peer-Durchschnitt", ""]
    if ins["schwaechen"]:
        lines += ["| Kennzahl | Zielbank | Peer-Ø | Abweichung | Rang |", "|---|---:|---:|---:|---:|"]
        for it in ins["schwaechen"][:15]:
            art = " *(berechnet)*" if it["berechnet"] else ""
            lines.append(f"| {it['label']}{art} | {_fmt(it['wert'])} {it['unit']} | "
                         f"{_fmt(it['peerAvg'])} | "
                         f"{'+' if (it['abweichungPct'] or 0) >= 0 else ''}"
                         f"{it['abweichungPct']} % | {it['rang'] or '–'} |")
    else:
        lines.append("Keine auswertbaren Kennzahlen unter dem Peer-Durchschnitt.")

    lines += ["", "## Kennzahlenmatrix", "",
              "| Kennzahl | Einheit | " + " | ".join(b["name"] for b in banks) + " |",
              "|---|---|" + "---:|" * len(banks)]
    for key in _kpi_keys(include_derived):
        d = DEFS[key]
        vals = [(mat.get(b["id"]) or {}).get(year, {}).get(key, {}).get("value")
                for b in banks]
        if all(v is None for v in vals):
            continue
        art = " *(ber.)*" if d.get("berechnet") else ""
        lines.append(f"| {d['label']}{art} | {d['unit']} | "
                     + " | ".join("–" if v is None else str(_fmt(v)) for v in vals) + " |")

    lines += ["", "## Untersuchungsbereiche", ""]
    for i, ub in enumerate(analytics.bereiche_status(data, year, bank_ids), start=1):
        lines += [f"### {i}. {ub['titel']}", "", ub["beschreibung"], "",
                  f"Datenabdeckung in der Vergleichsgruppe: **{ub['abdeckung']} %**", "",
                  f"*Datenlücke:* {ub['luecke']}", ""]

    lines += ["## Methodische Einschränkungen", ""]
    for item in data.get("methodik", {}).get("einschraenkungen", []):
        lines.append(f"- {item}")

    return "\n".join(lines) + "\n"
